from typing import Dict

from docx import Document
from docx.shared import Pt

from ..context import ResumeRenderContext
from .common import replace_placeholders_in_doc


def render_header(doc: Document, ctx: ResumeRenderContext) -> None:
    field_values: Dict[str, str] = {
        "name": ctx.name or "",
        "email": ctx.email or "",
        "github": ctx.github or "",
        "linkedin": ctx.linkedin or "",
        "portfolio": ctx.portfolio or "",
        "phone": ctx.phone or "",
        "location": ctx.location or "",
        "header_line": ctx.header_line or "",
    }
    replace_placeholders_in_doc(doc, field_values)
    _ensure_header_line_spacing(doc)


def _ensure_header_line_spacing(doc: Document) -> None:
    target_keys = {"{header_line}", "{ header_line }"}
    for para in doc.paragraphs:
        text = para.text or ""
        if any(k in text for k in target_keys):
            if para.paragraph_format and para.paragraph_format.space_after is None:
                para.paragraph_format.space_after = Pt(6)  # ~0.08"

